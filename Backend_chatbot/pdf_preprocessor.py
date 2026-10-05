import os
import re
from glob import glob

import fitz  # PyMuPDF — faster and more reliable than pdfplumber
from tqdm import tqdm
from unidecode import unidecode

# OCR support for scanned (image-based) PDFs
try:
    import pytesseract
    from pdf2image import convert_from_path
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False
    print("⚠️  OCR not available: install pytesseract and pdf2image for scanned PDF support")

# Word (.docx) support. Imported lazily-but-once so a missing dependency degrades
# to "Word uploads unsupported" instead of breaking PDF ingestion.
try:
    import docx  # python-docx
    from docx.table import Table as _DocxTable
    from docx.text.paragraph import Paragraph as _DocxParagraph
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("⚠️  Word support not available: install python-docx for .docx ingestion")

# OCR of images embedded in a .docx needs Pillow (pytesseract reads PIL images).
try:
    from PIL import Image as _PILImage
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

PDF_DIR = "pdfs"  # Changed from "books" to match your folder
OUTPUT_DIR = "data/txts"  # Changed to match your pipeline
COMBINED_FILE = "combined_book.txt"

# Every document type the ingestion pipeline can read. Extensions are compared
# lowercased. '.doc' (legacy binary Word) is deliberately NOT here — python-docx
# cannot read it; callers surface a "re-save as .docx" message instead.
SUPPORTED_EXTENSIONS = (".pdf", ".docx")

os.makedirs(OUTPUT_DIR, exist_ok=True)


# ---------- SMALL HELPERS ----------

def is_page_number(line: str) -> bool:
    """Detect plain or simple 'Page X' style page numbers."""
    stripped = line.strip()
    if not stripped:
        return False
    # Only digits
    if re.fullmatch(r"\d{1,4}", stripped):
        return True
    # Page 12, P. 12, etc.
    if re.fullmatch(r"(Page|PAGE|Pg\.?|PG\.?)\s*\d{1,4}", stripped):
        return True
    return False


def is_horizontal_rule(line: str) -> bool:
    """Lines made mostly of hyphens/underscores/equal signs."""
    stripped = line.strip()
    if not stripped:
        return False
    return bool(re.fullmatch(r"[-_=]{4,}", stripped))


def is_scanning_artifact(line: str) -> bool:
    """Isolated junk like '—', '·', or tiny symbol-only lines."""
    stripped = line.strip()
    # single long dash or bullet-only line
    if stripped in {"—", "–", "·", "•"}:
        return True
    # very short symbol-only strings (2–3 chars) with no letters/digits
    if len(stripped) <= 3 and not re.search(r"[A-Za-z0-9]", stripped):
        return True
    return False


def is_continued_on_next_page(line: str) -> bool:
    stripped = line.strip().lower()
    return "continued on next page" in stripped or "contd. on next page" in stripped


def normalize_bullets(line: str) -> str:
    """Convert various bullet symbols + indentation to simple hyphen bullets."""
    # Replace usual bullet symbols with hyphen
    line = re.sub(r"^[ \t]*[•·◦▪●►]+[ \t]*", "- ", line)

    # Book-specific: 'y ' used as bullet → convert to '- '
    line = re.sub(r"^\s*y\s+", "- ", line)

    # Normalize numbered lists like "1)" or "1." at start → keep as is, ensure space
    line = re.sub(r"^([ \t]*\d+[\).])\s*", r"\1 ", line)

    return line


def is_heading_candidate(line: str) -> bool:
    """
    CRITICAL: Enhanced heading detection for proper hierarchy.
    This is essential for your RAG pipeline to work correctly.
    
    FIX for Issue #9: Tightened Pattern 5 (Title Case) to require that
    MOST words are capitalized (true Title Case), not just the first word.
    This prevents regular sentences from being misclassified as headings.
    """
    stripped = line.strip()
    if not stripped:
        return False
    
    # PATTERN 1: Chapter/Unit headers (CRITICAL FOR HIERARCHY)
    if re.match(r"^(Chapter|CHAPTER|Unit|UNIT|CH\.?)\s+\d+", stripped, re.IGNORECASE):
        return True
    
    # PATTERN 2: Numbered sections like "1.2.3 Title"
    if re.match(r"^\d+(\.\d+)*\s+[A-Z]", stripped):
        return True
    
    # PATTERN 3: Lettered subsections like "A. Title"
    if re.match(r"^[A-Z]\.\s+[A-Z]", stripped):
        return True
    
    # PATTERN 4: ALL CAPS headers (but not too long)
    if stripped.isupper() and 3 < len(stripped) < 80 and not stripped.endswith("."):
        return True
    
    # PATTERN 5 (FIXED for Issue #9): Title Case headers
    # Must be short AND most words must be capitalized (true Title Case)
    # This prevents lines like "Media has changed over time" from being
    # misclassified as headings.
    if (len(stripped) < 80 and 
        not stripped.endswith((".", "?", "!", ":")) and
        stripped.count(" ") <= 8):
        words = stripped.split()
        if len(words) >= 2:
            # Skip small words for Title Case check
            small_words = {'a', 'an', 'the', 'and', 'or', 'but', 'in', 'on', 
                          'at', 'to', 'for', 'of', 'with', 'by', 'is', 'as'}
            significant_words = [w for w in words if w.lower() not in small_words]
            if significant_words:
                capitalized = sum(1 for w in significant_words if w[0].isupper())
                ratio = capitalized / len(significant_words)
                # At least 80% of significant words must be capitalized
                if ratio >= 0.8:
                    return True
    
    return False


def looks_like_table_row(line: str) -> bool:
    """
    Detect a table row in extracted text.
    NOTE: Tables will be converted to bullet format for better RAG retrieval.
    """
    if "  " not in line:
        return False
    if not re.search(r"\d", line):
        return False
    if len(line.strip()) < 10:
        return False
    return True


def split_table_row(line: str):
    """Split a table row on 2+ spaces."""
    parts = re.split(r"\s{2,}", line.strip())
    return [p.strip() for p in parts if p.strip()]


def convert_table_block(block_lines):
    """
    Convert a block of table-like lines to bullet format for better RAG.
    Format:
    Table: header1, header2, header3
    - row1_col1: col2=val2, col3=val3
    """
    if not block_lines:
        return []

    # First line = header row
    header_row = split_table_row(block_lines[0])
    if len(header_row) < 2:
        # Not a good table → return as-is
        return block_lines

    bullets = []
    table_title = "Table: " + ", ".join(header_row)
    bullets.append(table_title)

    for row in block_lines[1:]:
        cols = split_table_row(row)
        if not cols:
            continue
        # pad/truncate to header length
        if len(cols) < len(header_row):
            cols += [""] * (len(header_row) - len(cols))
        cols = cols[:len(header_row)]

        key = cols[0]
        kv_pairs = []
        for h, v in zip(header_row[1:], cols[1:]):
            if v:
                kv_pairs.append(f"{h}={v}")
        if kv_pairs:
            bullets.append(f"- {key}: " + ", ".join(kv_pairs))
        else:
            bullets.append(f"- {key}")

    return bullets


# ---------- PAGE-LEVEL PROCESSING ----------

def preprocess_page_text(raw_page_text: str) -> list[str]:
    """
    Process a single page:
    - split into lines
    - remove page numbers, rules, scanning artifacts, continued markers
    - normalize bullets
    - PRESERVE HEADING STRUCTURE (critical!)
    """
    lines = raw_page_text.splitlines()
    cleaned_lines = []

    for line in lines:
        # Normalize unicode early
        line = unidecode(line)

        if is_page_number(line):
            continue
        if is_horizontal_rule(line):
            continue
        if is_scanning_artifact(line):
            continue
        if is_continued_on_next_page(line):
            continue

        # Remove ridiculous sequences of dots/underscores etc.
        line = re.sub(r"[._]{4,}", " ", line)

        # Normalize bullets / numbered lists
        line = normalize_bullets(line)

        # Strip trailing spaces, keep leading for now
        line = line.rstrip()

        cleaned_lines.append(line)

    return cleaned_lines


def remove_repeated_headers_footers(pages_lines: list[list[str]]) -> list[list[str]]:
    """
    Detect lines that repeat as first/last non-empty line on many pages
    and remove them (headers/footers).
    IMPORTANT: This preserves content headers like "Chapter 1" that only appear once.
    """
    first_lines = []
    last_lines = []

    for lines in pages_lines:
        non_empty = [l for l in lines if l.strip()]
        if not non_empty:
            continue
        first_lines.append(non_empty[0].strip())
        last_lines.append(non_empty[-1].strip())

    def get_repeated(candidates):
        counts = {}
        for l in candidates:
            counts[l] = counts.get(l, 0) + 1
        # Consider header/footer if appears on >= 3 pages
        # AND is not a content heading (Chapter, Unit, etc.)
        repeated = set()
        for l, c in counts.items():
            if c >= 3:
                # Don't remove if it's a chapter/unit/section header
                if not re.match(r"^(Chapter|Unit|CHAPTER|UNIT|\d+\.|[A-Z]\.)", l, re.IGNORECASE):
                    repeated.add(l)
        return repeated

    header_candidates = get_repeated(first_lines)
    footer_candidates = get_repeated(last_lines)

    cleaned_pages = []
    for lines in pages_lines:
        new_lines = []
        non_empty = [i for i, l in enumerate(lines) if l.strip()]
        header_idx = non_empty[0] if non_empty else None
        footer_idx = non_empty[-1] if non_empty else None

        for i, line in enumerate(lines):
            stripped = line.strip()
            if stripped and i == header_idx and stripped in header_candidates:
                continue
            if stripped and i == footer_idx and stripped in footer_candidates:
                continue
            new_lines.append(line)
        cleaned_pages.append(new_lines)

    return cleaned_pages


def detect_and_convert_tables(lines: list[str]) -> list[str]:
    """
    Scan through lines, group table-like blocks, convert them to bullets.
    This makes tables more RAG-friendly.
    """
    result = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]

        if looks_like_table_row(line):
            block = [line]
            i += 1
            while i < n and looks_like_table_row(lines[i]):
                block.append(lines[i])
                i += 1
            conv = convert_table_block(block)
            result.extend(conv)
        else:
            result.append(line)
            i += 1

    return result


def merge_lines_to_paragraphs(lines: list[str]) -> str:
    """
    CRITICAL FOR RAG: Merge lines into paragraphs while:
    - PRESERVING headings as separate lines (for hierarchy detection)
    - PRESERVING bullets / numbered lists as separate lines
    - Merging regular content into paragraphs
    """
    processed = []
    buffer = []

    def flush_buffer():
        if buffer:
            processed.append(" ".join(buffer).strip())
            buffer.clear()

    for line in lines:
        stripped = line.strip()

        if not stripped:
            # blank line → paragraph break
            flush_buffer()
            processed.append("")  # keep empty line
            continue

        is_bullet = re.match(r"^(-|\d+[\).])\s+", stripped) is not None
        heading = is_heading_candidate(line)

        if is_bullet or heading:
            # CRITICAL: Flush buffer and keep heading/bullet separate
            flush_buffer()
            processed.append(stripped)
        else:
            # Regular content - add to buffer for merging
            buffer.append(stripped)

    flush_buffer()

    # Collapse >2 empty lines to max 1 (but keep structure)
    final_lines = []
    empty_count = 0
    for l in processed:
        if l == "":
            empty_count += 1
            if empty_count <= 1:
                final_lines.append(l)
        else:
            empty_count = 0
            final_lines.append(l)

    return "\n".join(final_lines).strip()


def postprocess_global(text: str) -> str:
    """
    Global cleanups after paragraphs:
    - remove known running headers (CAREFULLY - don't remove content headers)
    - fix spaced section numbers
    - normalize spaces/newlines
    
    FIX for Issue #8: Made "& Videography" removal scoped to standalone
    running header lines only, instead of replacing it everywhere in content.
    """

    # 1) Remove running header "Digital Photography" ONLY if it's standalone
    #    Don't remove if it's part of a title like "Chapter 1: Digital Photography"
    text = re.sub(
        r"^Digital Photography\s*$",
        "",
        text,
        flags=re.MULTILINE
    )

    # 2) FIX for Issue #8: Only remove "& Videography" on standalone header lines
    #    e.g., "Digital Photography & Videography" as a running header
    #    Previously this removed "& Videography" from ALL content, corrupting
    #    sentences like "the course covers Photography & Videography techniques"
    text = re.sub(
        r"^Digital Photography\s*&\s*Videography\s*$",
        "",
        text,
        flags=re.MULTILINE
    )

    # 3) Fix patterns like "1. 2.1" → "1.2.1" (spaced section numbers)
    text = re.sub(r"(\d)\.\s+(\d)", r"\1.\2", text)

    # 4) Normalize spaces (but keep newlines for structure)
    text = re.sub(r"[ \t]+", " ", text)
    
    # 5) Collapse excessive newlines (max 2)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# ---------- MAIN PIPELINE ----------

def _ocr_page_image(pdf_path: str, page_index: int) -> str:
    """
    Render a single PDF page to an image and run Tesseract OCR on it.
    Used as a fallback for scanned / image-based pages.
    """
    if not OCR_AVAILABLE:
        return ""
    try:
        images = convert_from_path(
            pdf_path,
            dpi=300,               # High resolution for better OCR accuracy
            first_page=page_index + 1,
            last_page=page_index + 1,
        )
        if not images:
            return ""
        text = pytesseract.image_to_string(images[0], lang="eng")
        return text
    except Exception as e:
        print(f"⚠️  OCR failed on page {page_index + 1}: {e}")
        return ""


def extract_and_clean_pdf(pdf_path: str) -> str:
    """
    Main extraction pipeline - optimized for RAG with hierarchy preservation.
    Uses PyMuPDF (fitz) for fast, reliable text extraction (3-5x faster than pdfplumber).

    OCR Fallback: If a page returns no text (scanned / image-based PDF),
    it is automatically rendered as an image and processed with Tesseract OCR.
    This makes both digital PDFs and scanned PDFs work seamlessly.
    """
    pages_lines = []
    scanned_pages = 0

    doc = fitz.open(pdf_path)
    for page_index, page in enumerate(doc):
        raw = page.get_text() or ""

        # Fallback to OCR if page has no embedded text (scanned PDF)
        if not raw.strip() and OCR_AVAILABLE:
            print(f"   🔍 Page {page_index + 1}: no text found — running OCR...")
            raw = _ocr_page_image(pdf_path, page_index)
            if raw.strip():
                scanned_pages += 1

        page_lines = preprocess_page_text(raw)
        pages_lines.append(page_lines)
    doc.close()

    if scanned_pages > 0:
        print(f"   📷 OCR processed {scanned_pages} scanned page(s)")

    # Remove repeated headers / footers (preserves content headers)
    pages_lines = remove_repeated_headers_footers(pages_lines)

    # Flatten page lines with explicit page break (blank line)
    all_lines = []
    for pl in pages_lines:
        all_lines.extend(pl)
        all_lines.append("")  # Page break

    # Table detection & conversion to bullets
    all_lines = detect_and_convert_tables(all_lines)

    # Merge lines into paragraphs, PRESERVING bullets & headings
    cleaned_text = merge_lines_to_paragraphs(all_lines)

    # Global postprocessing
    cleaned_text = postprocess_global(cleaned_text)

    return cleaned_text


# ---------- WORD (.docx) EXTRACTION ----------

def _docx_iter_block_items(document):
    """
    Yield the document's paragraphs and tables in TRUE document order.

    python-docx exposes `.paragraphs` and `.tables` as two separate lists, which
    loses their relative order — a table appearing between two paragraphs would
    otherwise be appended at the end and break the reading flow (and therefore
    the section/heading hierarchy the chunker depends on). Walking the underlying
    body XML preserves the real order.
    """
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield _DocxParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield _DocxTable(child, document)


def _docx_paragraph_has_page_break(paragraph) -> bool:
    """True when the paragraph carries an explicit page break."""
    try:
        xml = paragraph._p.xml
        return 'w:br' in xml and 'type="page"' in xml
    except Exception:
        return False


def _docx_table_to_lines(table) -> list[str]:
    """
    Render a Word table directly into the SAME bullet format that
    convert_table_block() produces for PDF tables:

        Table: Medium, Reach, Year
        - Radio: Reach=Wide rural reach, Year=1927

    We emit the final form rather than re-serialising to whitespace-separated
    text and letting the PDF heuristic re-detect it. That heuristic
    (looks_like_table_row) requires a digit in the row, so an all-text header
    like "Medium  Reach  Year" would be missed — the first DATA row would then
    be mistaken for the header and every cell would be labelled with the wrong
    column name. Word already gives us the true structure, so we use it.

    The emitted lines contain no double spaces, so detect_and_convert_tables()
    correctly leaves them alone downstream.
    """
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cells.append(" ".join(cell.text.split()))
        # Collapse horizontally-merged cells (python-docx repeats their text).
        deduped = []
        for c in cells:
            if not deduped or c != deduped[-1]:
                deduped.append(c)
        if any(c for c in deduped):
            rows.append(deduped)

    if not rows:
        return []

    header = rows[0]
    # Single-row table, or a header that is really just prose: emit plainly.
    if len(rows) == 1 or len([h for h in header if h]) < 2:
        return [", ".join(c for c in r if c) for r in rows]

    lines = ["Table: " + ", ".join(h for h in header if h)]
    for r in rows[1:]:
        padded = r + [""] * (len(header) - len(r))
        key = padded[0] or "(row)"
        pairs = [
            f"{h}={v}"
            for h, v in zip(header[1:], padded[1:])
            if h and v
        ]
        lines.append(f"- {key}: " + ", ".join(pairs) if pairs else f"- {key}")
    return lines


def _ocr_docx_images(docx_path: str) -> str:
    """
    OCR every image embedded in a .docx.

    Mirrors the PDF OCR fallback: a Word file whose content is pasted scans /
    screenshots has no extractable paragraph text, so without this the document
    would ingest as empty. Images live under `word/media/` inside the .docx zip.
    """
    if not (OCR_AVAILABLE and PIL_AVAILABLE):
        return ""

    import io
    import zipfile

    chunks = []
    try:
        with zipfile.ZipFile(docx_path) as zf:
            media = [n for n in zf.namelist() if n.startswith("word/media/")]
            for name in sorted(media):
                ext = os.path.splitext(name)[1].lower()
                if ext not in (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".gif"):
                    continue
                try:
                    with zf.open(name) as fh:
                        image = _PILImage.open(io.BytesIO(fh.read()))
                        image.load()
                    text = pytesseract.image_to_string(image, lang="eng")
                    if text.strip():
                        chunks.append(text)
                except Exception as e:
                    print(f"⚠️  OCR failed on embedded image {name}: {e}")
    except Exception as e:
        print(f"⚠️  Could not read embedded images from {docx_path}: {e}")

    return "\n".join(chunks)


def extract_and_clean_docx(docx_path: str) -> str:
    """
    Extract and clean a Word (.docx) document.

    Produces text in exactly the same shape as extract_and_clean_pdf(): the raw
    text acquisition differs, but every downstream cleaning stage
    (header/footer removal, table conversion, paragraph merging, global
    postprocessing) is shared, so Word and PDF content is indistinguishable to
    the chunker, BM25 index and retriever.

    OCR fallback: if the document yields no usable text, embedded images are
    OCR'd — the Word equivalent of a scanned PDF.
    """
    if not DOCX_AVAILABLE:
        raise RuntimeError(
            "Word support requires python-docx. Install it with: pip install python-docx"
        )

    document = docx.Document(docx_path)

    # Build pseudo-pages: Word has no fixed pagination, but explicit page breaks
    # give us natural boundaries. Keeping the page structure lets
    # remove_repeated_headers_footers() work the same way it does for PDFs.
    pages_raw: list[list[str]] = [[]]
    for block in _docx_iter_block_items(document):
        if isinstance(block, _DocxTable):
            pages_raw[-1].extend(_docx_table_to_lines(block))
            continue

        if _docx_paragraph_has_page_break(block) and pages_raw[-1]:
            pages_raw.append([])

        text = block.text.rstrip()
        # Word list items lose their bullet glyph in .text; re-add one so
        # normalize_bullets()/merge_lines_to_paragraphs() keep them as list items
        # instead of merging them into a paragraph.
        style = (getattr(block.style, "name", "") or "").lower()
        if text.strip() and ("list" in style) and not re.match(r"^\s*([-•*]|\d+[.)])\s", text):
            text = f"- {text.strip()}"
        pages_raw[-1].append(text)

    # Section headers/footers repeat on every page in Word; the PDF path strips
    # such repetition, so we simply do not import them.

    total_chars = sum(len(l.strip()) for p in pages_raw for l in p)
    if total_chars < 20:
        print("   🔍 No text found in Word document — running OCR on embedded images...")
        ocr_text = _ocr_docx_images(docx_path)
        if ocr_text.strip():
            print("   📷 OCR recovered text from embedded image(s)")
            pages_raw = [ocr_text.split("\n")]

    # From here on: identical to the PDF pipeline.
    pages_lines = [preprocess_page_text("\n".join(page)) for page in pages_raw]
    pages_lines = remove_repeated_headers_footers(pages_lines)

    all_lines = []
    for pl in pages_lines:
        all_lines.extend(pl)
        all_lines.append("")  # Page break

    all_lines = detect_and_convert_tables(all_lines)
    cleaned_text = merge_lines_to_paragraphs(all_lines)
    cleaned_text = postprocess_global(cleaned_text)

    return cleaned_text


# ---------- FORMAT-AGNOSTIC ENTRY POINTS ----------

def is_supported_document(path: str) -> bool:
    """True when the file extension is one the pipeline can ingest."""
    return os.path.splitext(path)[1].lower() in SUPPORTED_EXTENSIONS


def extract_and_clean_document(path: str) -> str:
    """
    Extract and clean ANY supported document (.pdf or .docx).

    This is the entry point ingestion code should call; it dispatches on the file
    extension so callers never need to care which format they were given.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_and_clean_pdf(path)
    if ext == ".docx":
        return extract_and_clean_docx(path)
    if ext == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported. Please re-save the file as .docx "
            "(Word: File > Save As > Word Document (.docx)) and upload again."
        )
    raise ValueError(
        f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
    )


def count_document_pages(path: str) -> int:
    """
    Best-effort page count, used only for progress reporting.

    PDFs report their true page count. Word has no fixed pagination until it is
    rendered, so we count explicit page breaks (+1); a document with none counts
    as a single page.
    """
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            doc = fitz.open(path)
            n = doc.page_count
            doc.close()
            return n
        if ext == ".docx" and DOCX_AVAILABLE:
            document = docx.Document(path)
            breaks = sum(
                1 for p in document.paragraphs if _docx_paragraph_has_page_break(p)
            )
            return breaks + 1
    except Exception as e:
        print(f"⚠️  Could not count pages for {path}: {e}")
    return 1


def process_all_pdfs():
    """
    Process all PDFs and create individual + combined files.
    """
    # Pick up every supported document type, not just PDFs.
    pdf_paths = []
    for ext in SUPPORTED_EXTENSIONS:
        pdf_paths.extend(glob(os.path.join(PDF_DIR, f"*{ext}")))
    # Skip Word lock/temp files (~$name.docx) that Word leaves behind while open.
    pdf_paths = sorted(p for p in pdf_paths if not os.path.basename(p).startswith("~$"))
    if not pdf_paths:
        print(f"❌ No documents found in folder: {PDF_DIR}")
        print(f"   Supported types: {', '.join(SUPPORTED_EXTENSIONS)}")
        return

    print(f"\n{'='*70}")
    print(f"📚 Found {len(pdf_paths)} document(s) to process")
    print(f"{'='*70}")

    combined_texts = []

    for pdf_path in tqdm(pdf_paths, desc="Processing documents"):
        filename = os.path.basename(pdf_path)
        print(f"\n{'='*70}")
        print(f"📖 Processing: {filename}")
        print(f"{'='*70}")

        try:
            cleaned_text = extract_and_clean_document(pdf_path)

            # Save individual file
            out_path = os.path.join(
                OUTPUT_DIR,
                os.path.splitext(filename)[0] + ".txt"
            )
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(cleaned_text)
            
            print(f"✅ Saved to: {out_path}")
            print(f"   Length: {len(cleaned_text)} characters")

            # Add to combined with clear separator
            combined_texts.append(f"\n\n{'='*70}\n=== SOURCE: {filename} ===\n{'='*70}\n\n{cleaned_text}")

        except Exception as e:
            print(f"❌ Error processing {filename}: {e}")
            import traceback
            traceback.print_exc()

    # Save combined file
    combined_path = os.path.join(OUTPUT_DIR, COMBINED_FILE)
    with open(combined_path, "w", encoding="utf-8") as f:
        f.write("\n".join(combined_texts))

    print(f"\n{'='*70}")
    print("✅ PROCESSING COMPLETE!")
    print(f"{'='*70}")
    print(f"📁 Individual files: {OUTPUT_DIR}")
    print(f"📄 Combined file: {combined_path}")
    print(f"\n🔄 Next steps:")
    print(f"   1. Review the output files to verify structure is preserved")
    print(f"   2. Look for 'Chapter', 'Unit', numbered sections in the output")
    print(f"   3. Run: python process_txt_pipeline.py")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    process_all_pdfs()
